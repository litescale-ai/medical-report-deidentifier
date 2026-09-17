"""Real file regressions for extraction, same-format replacement and restoration."""

import asyncio
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from agents.transcriber import transcribe_media
from utils.document_formats import extract_document
from utils.document_editor import deidentify_document, reidentify_document

PATIENT = "Alex Example"
DOCTOR = "Robin Tester"
TEXT = f"Patient: {PATIENT}.\nDoctor: {DOCTOR}.\nAssessment: improving mobility.\n"


def create_fixtures(directory):
    """Create small synthetic records with format-specific secondary text surfaces."""
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from pptx import Presentation
    from pptx.util import Inches
    import pymupdf

    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "record.txt").write_text(TEXT)
    (directory / "record.md").write_text("# Assessment\n\n" + TEXT)
    (directory / "record.html").write_text(
        '<!doctype html><html><body><h1>Assessment</h1><p>Patient: Alex <b>Example</b>.</p>'
        '<p title="Alex&#32;Example">Doctor: Robin Tester.</p><p>Assessment: improving mobility.</p></body></html>')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assessment"
    sheet["A1"] = f"Patient: {PATIENT}"
    sheet["A2"] = f"Doctor: {DOCTOR}"
    sheet["A3"] = "Assessment: improving mobility."
    sheet["A1"].font = __import__('openpyxl').styles.Font(bold=True)
    sheet["A1"].comment = Comment(f"Seen by {DOCTOR}", DOCTOR)
    sheet["B1"] = 42
    sheet["B2"] = "=B1+1"
    workbook.save(directory / "record.xlsx")
    workbook.close()

    document = Document()
    paragraph = document.add_paragraph("Patient: ")
    paragraph.add_run("Alex ").bold = True
    paragraph.add_run("Example")
    paragraph.add_run("; Alex ")
    paragraph.add_run("Example")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    link_paragraph = document.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "patient")
    link_run = OxmlElement("w:r")
    link_text = OxmlElement("w:t")
    link_text.text = PATIENT
    link_run.append(link_text)
    hyperlink.append(link_run)
    link_paragraph._p.append(hyperlink)
    document.add_paragraph("Assessment: improving mobility.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = f"Doctor: {DOCTOR}"
    document.sections[0].header.paragraphs[0].text = f"Patient: {PATIENT}"
    document.sections[0].footer.paragraphs[0].text = f"Doctor: {DOCTOR}"
    document.tables[0].cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = PATIENT
    document.save(directory / "record.docx")

    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((40, 60), TEXT)
        document.save(directory / "record.pdf")

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Patient: Alex "
    run.font.bold = True
    paragraph.add_run().text = "Example"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1)).text_frame.text = "Assessment: improving mobility."
    slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(7), Inches(1)).table.cell(0, 0).text = f"Doctor: {DOCTOR}"
    slide.notes_slide.notes_text_frame.text = f"Seen by {DOCTOR}; patient {PATIENT}."
    deck.save(directory / "record.pptx")
    return sorted(directory.glob("record.*"))


def extracted_text(path):
    return "\n".join(location + "\n" + text for location, text in extract_document(path))


class DocumentFormatsTest(unittest.IsolatedAsyncioTestCase):
    async def test_each_format_extracts_replaces_and_restores_real_files(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            for source in create_fixtures(Path(directory) / "input"):
                with self.subTest(format=source.suffix):
                    with patch("agents.transcriber.Agent", side_effect=AssertionError("Model used for local extraction")):
                        result = await transcribe_media(str(source), backend="ollama")
                    content = "\n".join(item["content"] for item in result["items"])
                    self.assertIn(PATIENT, content)
                    self.assertIn(DOCTOR, content)
                    output = Path(directory) / ("redacted" + source.suffix)
                    # The shorter alias must not interrupt a full-name match spanning runs.
                    replacements = {PATIENT: "PATIENT_TEST", DOCTOR: "DOCTOR_TEST", "Alex": "PATIENT_TEST"}
                    self.assertTrue(deidentify_document(str(source), str(output), replacements))
                    text = extracted_text(output)
                    self.assertNotIn(PATIENT, text)
                    self.assertNotIn(DOCTOR, text)
                    self.assertNotIn("Example", text)
                    self.assertIn("PATIENT_TEST", text)
                    self.assertIn("DOCTOR_TEST", text)
                    self.assertIn("improving mobility", text)
                    restored = Path(directory) / ("restored" + source.suffix)
                    self.assertTrue(reidentify_document(str(output), str(restored), {
                        "PATIENT_TEST": {"canonical_name": PATIENT}, "DOCTOR_TEST": {"canonical_name": DOCTOR}}))
                    text = extracted_text(restored)
                    self.assertIn(PATIENT, text)
                    self.assertIn(DOCTOR, text)
                    self.assertNotIn("PATIENT_TEST", text)
                    if source.suffix == ".xlsx":
                        workbook = load_workbook(output)
                        self.assertEqual(workbook.active["B1"].value, 42)
                        self.assertEqual(workbook.active["B2"].value, "=B1+1")
                        self.assertTrue(workbook.active["A1"].font.bold)
                        self.assertNotIn(DOCTOR, workbook.active["A1"].comment.text)
                        workbook.close()
                    elif source.suffix == ".html":
                        self.assertIn("<b>", output.read_text())

    async def test_media_timeout_names_stage_and_file(self):
        class StalledAgent:
            def __init__(self, **kwargs):
                pass
            async def __aenter__(self):
                return self
            async def __aexit__(self, *args):
                return False
            async def chat(self, _):
                await asyncio.Event().wait()
        with patch("agents.transcriber.load_media_file", return_value=object()), patch(
            "agents.transcriber.Agent", StalledAgent
        ), patch.dict("os.environ", {"MODEL_TIMEOUT_SECONDS": "0.01"}):
            with self.assertRaisesRegex(TimeoutError, "Transcription of interview.wav exceeded"):
                await transcribe_media("interview.wav", backend="ollama")

    def test_spreadsheet_formula_tokens_and_sheet_references(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.xlsx", Path(directory) / "output.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = PATIENT
            sheet["B1"] = 5
            sheet["A1"] = '=MAX(B1:B10)'
            sheet["A2"] = "='Alex Example'!B1"
            sheet["A3"] = '=IF(B1=5,"Alex Example","MAX")'
            workbook.save(source)
            deidentify_document(str(source), str(output), {PATIENT: "PATIENT_TEST", "MAX": "PATIENT_MAX"})
            result = load_workbook(output)
            self.assertEqual(result.active.title, "PATIENT_TEST")
            self.assertEqual(result.active["A1"].value, '=MAX(B1:B10)')
            self.assertEqual(result.active["A2"].value, "='PATIENT_TEST'!B1")
            self.assertEqual(result.active["A3"].value, '=IF(B1=5,"PATIENT_TEST","PATIENT_MAX")')
            result.close()

    def test_html_preserves_scripts_and_replaces_encoded_attributes(self):
        from utils.document_formats import HtmlText
        source = '<script>if (a < b && c > d) {run();}</script><p title="Alex&#32;Example">Alex Example</p>'
        result = HtmlText(source).replace({PATIENT: "PATIENT_TEST"})
        self.assertIn('if (a < b && c > d) {run();}', result)
        self.assertIn('title="PATIENT_TEST"', result)

    def test_timeout_rejects_nonfinite_and_invalid_values(self):
        from utils.agent_config import model_timeout_seconds
        for value in ("nan", "inf", "-inf", "0", "-1", "invalid"):
            with self.subTest(value=value), patch.dict("os.environ", {"MODEL_TIMEOUT_SECONDS": value}):
                with self.assertRaisesRegex(ValueError, "finite positive"):
                    model_timeout_seconds()

    async def test_legacy_office_formats_require_conversion(self):
        for extension in ("doc", "xls", "ppt"):
            with self.subTest(format=extension), self.assertRaisesRegex(ValueError, "Convert"):
                await transcribe_media("record." + extension, backend="ollama")

    async def test_unicode_text_extensions(self):
        with tempfile.TemporaryDirectory() as directory:
            for extension in (".txt", ".md", ".html", ".csv", ".yaml", ".xml"):
                path = Path(directory) / ("record" + extension)
                path.write_text("Patient: Zoë Example", encoding="utf-16")
                result = await transcribe_media(str(path), backend="ollama")
                self.assertIn("Zoë Example", result["items"][0]["content"])


if __name__ == "__main__":
    unittest.main()
