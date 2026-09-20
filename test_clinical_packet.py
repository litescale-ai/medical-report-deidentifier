"""Offline safety regressions for the experiment; model quality is measured separately."""
import asyncio
from contextlib import redirect_stderr, redirect_stdout
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
import unittest

import prepare_for_ai
from utils.clinical_packet import NerDetector, Section, extract_case, ner_windows, prepare_packet, write_packet


class Detector:
    metadata = {'detector': 'test'}
    def __init__(self, entities): self.entities = entities
    async def discover(self, text): return self.entities


def person(name, aliases=()):
    return {'canonical_name': name, 'entity_type': 'PERSON', 'variations': list(aliases)}


class PacketTests(unittest.TestCase):
    def test_cli_defaults_and_saved_additions_apply_to_later_packets(self):
        with TemporaryDirectory() as root:
            root = Path(root)
            saved = root / 'private/keep.txt'
            source = root / 'report.txt'
            terms = [*prepare_for_ai.DEFAULT_KEEP_TERMS, 'Digit span']
            source.write_text('Patient: Thandi Dlamini.\n' + '\n'.join(terms))
            entities = [person('Thandi Dlamini'), *[person(term) for term in terms]]
            with patch.object(prepare_for_ai, 'SAVED_KEEP_TERMS', saved), redirect_stdout(io.StringIO()):
                with patch.object(prepare_for_ai, 'NerDetector') as model:
                    prepare_for_ai.cli(['--list-keep-terms'])
                    self.assertFalse(saved.exists())
                    prepare_for_ai.cli(['--add-keep-term', 'Digit span', '--add-keep-term', '  DIGIT   SPAN  '])
                    model.assert_not_called()
                self.assertEqual(saved.read_text().splitlines(), terms)
                self.assertEqual(saved.stat().st_mode & 0o777, 0o600)
                with patch.object(prepare_for_ai, 'NerDetector', return_value=Detector(entities)):
                    prepare_for_ai.cli([str(source), '--output', str(root / 'out')])
                manifest = json.loads((root / 'out/PRIVATE.json').read_text())
                packet = (root / 'out/REVIEW_REQUIRED.md').read_text()
                self.assertEqual(manifest['keep_terms'], terms)
                for term in terms:
                    self.assertIn(term, packet)
                self.assertNotIn('Thandi Dlamini', packet)
                # Direct edits, including removals, are respected on the next run.
                saved.write_text('Digit span\n')
                with patch.object(prepare_for_ai, 'NerDetector', return_value=Detector(entities)):
                    prepare_for_ai.cli([str(source), '--output', str(root / 'edited')])
                edited = (root / 'edited/REVIEW_REQUIRED.md').read_text()
                self.assertIn('Digit span', edited)
                self.assertNotIn('timed classroom', edited)

    def test_cli_custom_list_and_disabled_terms_do_not_mutate_saved_list(self):
        with TemporaryDirectory() as root:
            root = Path(root)
            saved, custom, source = root / 'saved.txt', root / 'custom.txt', root / 'report.txt'
            saved.write_text('Beery VMI\n')
            custom.write_text('DAP\n')
            source.write_text('Beery VMI and DAP')
            with patch.object(prepare_for_ai, 'SAVED_KEEP_TERMS', saved), \
                 patch.object(prepare_for_ai, 'NerDetector', return_value=Detector([person('Beery VMI'), person('DAP')])), \
                 redirect_stdout(io.StringIO()):
                prepare_for_ai.cli([str(source), '--keep-terms', str(custom), '--output', str(root / 'custom')])
                prepare_for_ai.cli([str(source), '--no-keep-terms', '--output', str(root / 'disabled')])
            packet = (root / 'custom/REVIEW_REQUIRED.md').read_text()
            self.assertIn('DAP', packet)
            self.assertNotIn('Beery VMI', packet)
            disabled = json.loads((root / 'disabled/PRIVATE.json').read_text())
            self.assertEqual(disabled['keep_terms'], [])
            self.assertEqual(saved.read_text(), 'Beery VMI\n')

    def test_cli_first_run_seeds_the_editable_list(self):
        with TemporaryDirectory() as root:
            root = Path(root)
            saved, source = root / 'keep.txt', root / 'report.txt'
            source.write_text('Dose 5 mg.')
            with patch.object(prepare_for_ai, 'SAVED_KEEP_TERMS', saved), \
                 patch.object(prepare_for_ai, 'NerDetector', return_value=Detector([])), \
                 redirect_stdout(io.StringIO()):
                prepare_for_ai.cli([str(source), '--output', str(root / 'out')])
            self.assertEqual(saved.read_text().splitlines(), list(prepare_for_ai.DEFAULT_KEEP_TERMS))

    def test_cli_invalid_arguments_do_not_write_settings_or_load_model(self):
        with TemporaryDirectory() as root:
            saved = Path(root) / 'keep.txt'
            with patch.object(prepare_for_ai, 'SAVED_KEEP_TERMS', saved), \
                 patch.object(prepare_for_ai, 'NerDetector') as model, redirect_stderr(io.StringIO()):
                for args in [[], ['report.pdf'], ['--add-keep-term', ' '],
                             ['--no-keep-terms', '--add-keep-term', 'DAP'],
                             ['--output', 'unused', '--list-keep-terms'],
                             ['--keep-terms', str(Path(root) / 'missing.txt'), '--list-keep-terms'],
                             ['report.pdf', '--output', root, '--add-keep-term', 'DAP']]:
                    with self.subTest(args=args), self.assertRaises(SystemExit):
                        prepare_for_ai.cli(args)
                model.assert_not_called()
                self.assertFalse(saved.exists())

    def test_case_replacements_rules_and_clinical_preservation(self):
        sections = [Section('Document 1', 'Patient: Thandi Dlamini\nID number: 1203155009084\n'
                            'Email: thandi@example.test\nDose 5 mg. Blood pressure 120/80.'),
                    Section('Document 2', 'THANDI  DLAMINI reports improved sleep. Assessment: 2026-08-12.')]
        result = asyncio.run(prepare_packet(sections, Detector([person('Thandi Dlamini')])))
        self.assertNotIn('Dlamini', result['markdown'])
        self.assertNotIn('1203155009084', result['markdown'])
        self.assertNotIn('thandi@example.test', result['markdown'])
        self.assertEqual(result['markdown'].count('[PERSON_1]'), 2)
        self.assertIn('Dose 5 mg. Blood pressure 120/80.', result['markdown'])
        self.assertIn('Assessment: 2026-08-12.', result['markdown'])
        with TemporaryDirectory() as root:
            out = Path(root)/'packet'
            write_packet(result, out, ['private-patient-name.pdf'])
            self.assertEqual((out/'PRIVATE.json').stat().st_mode & 0o777, 0o600)
            self.assertNotIn('private-patient-name.pdf', (out/'REVIEW_REQUIRED.md').read_text())
            with self.assertRaises(FileExistsError): write_packet(result, out, [])

    def test_invalid_or_conflicting_entities_fail_closed(self):
        sections = [Section('Document 1', 'Patient: Lee Jacobs. Lee is here. Other: Lee Adams.')]
        for entities in [[person('Hallucinated Name')],
                         [person('Lee Jacobs', ['Lee']), person('Lee Adams', ['Lee'])]]:
            with self.subTest(entities=entities), self.assertRaises(ValueError):
                asyncio.run(prepare_packet(sections, Detector(entities)))

    def test_role_aliases_do_not_replace_clinical_prose(self):
        result = asyncio.run(prepare_packet(
            [Section('Document 1', 'Patient: Thandi Dlamini. The patient reports improved sleep.')],
            Detector([person('Thandi Dlamini', ['Patient', 'the patient'])])))
        self.assertIn('Patient: [PERSON_1]. The patient reports improved sleep.', result['markdown'])

    def test_invalid_ner_offsets_fail_instead_of_redacting_wrong_text(self):
        detector = NerDetector.__new__(NerDetector)
        detector.threshold = .3
        detector.model = SimpleNamespace(data_processor=None, predict_entities=lambda *a, **k: [
            {'start': 0, 'end': 5, 'text': 'Wrong', 'label': 'name'}])
        with patch('utils.clinical_packet.ner_windows', return_value=['Alice has a score of 85.']):
            with self.assertRaisesRegex(ValueError, 'invalid span'):
                asyncio.run(detector.discover('Alice has a score of 85.'))

    def test_reviewed_terms_preserve_content_while_people_are_redacted(self):
        keep = ['timed classroom', 'SNAP-IV Rating Scales', 'FCPaed', 'Beery VMI',
                'Beery-Buktenica', 'VMI', 'DAP', 'Kaleidovision', 'technology startup', 'Grade 2']
        text = 'Patient: Thandi Dlamini. School: Example Primary School.\n' + '\n'.join(keep)
        entities = [person('Thandi Dlamini'),
                    {'canonical_name': 'Example Primary School', 'entity_type': 'ORGANIZATION', 'variations': []}]
        entities += [{'canonical_name': value, 'entity_type': 'ORGANIZATION', 'variations': []} for value in keep]
        result = asyncio.run(prepare_packet([Section('Document 1', text)], Detector(entities), keep_terms=keep))
        for value in keep:
            with self.subTest(value=value):
                self.assertIn(value, result['markdown'])
        self.assertNotIn('Thandi Dlamini', result['markdown'])
        self.assertNotIn('Example Primary School', result['markdown'])
        mapped = [value for values in result['mapping'].values() for value in values]
        self.assertTrue(all(value not in mapped for value in keep))
        self.assertEqual(result['keep_terms'], keep)

    def test_keep_terms_protect_partial_aliases_but_not_nearby_identifiers(self):
        text = 'SNAP-IV\nRating Scales completed by Sam Example.\nClinical term: timed classroom.'
        entities = [person('Sam Example'), person('SNAP-IV'), person('Scales'), person('timed classroom')]
        result = asyncio.run(prepare_packet([Section('Document 1', text)], Detector(entities),
                                           keep_terms=['snap-iv rating scales', 'TIMED CLASSROOM']))
        self.assertIn('SNAP-IV\nRating Scales', result['markdown'])
        self.assertIn('timed classroom', result['markdown'])
        self.assertNotIn('Sam Example', result['markdown'])
        self.assertEqual(len(result['mapping']), 1)
        with self.assertRaisesRegex(ValueError, 'crosses a reviewed'):
            asyncio.run(prepare_packet([Section('Document 1', 'Example Primary School')],
                        Detector([person('Example Primary School')]), keep_terms=['School']))

    def test_source_markup_is_literal_and_cannot_close_fence(self):
        text='```\n![scan](https://example.test/private)\n<script>fetch("https://example.test")</script>'
        result=asyncio.run(prepare_packet([Section('Document 1', text)], Detector([])))
        self.assertIn('````text\n'+text+'\n````', result['markdown'])

    def test_docx_table_order_and_private_headers(self):
        from docx import Document
        with TemporaryDirectory() as root:
            path=Path(root)/'Secret Name.docx'
            doc=Document()
            doc.add_paragraph('Before')
            t=doc.add_table(rows=1,cols=2)
            t.cell(0,0).text='Memory'; t.cell(0,1).text='85'
            doc.add_paragraph('After')
            doc.sections[0].header.paragraphs[0].text='Secret Name'
            doc.save(path)
            sections=extract_case([path])
            self.assertIn('Before\nMemory | 85\nAfter', sections[0].text)
            result=asyncio.run(prepare_packet(sections, Detector([person('Secret Name')])))
            self.assertNotIn('Secret Name', result['markdown'])
            self.assertIn('Memory | 85', result['markdown'])

    def test_empty_or_unsupported_input_fails(self):
        with self.assertRaises(ValueError): extract_case([])
        with TemporaryDirectory() as root:
            p=Path(root)/'empty.txt'; p.write_text('')
            with self.assertRaises(ValueError): extract_case([p])

    def test_windows_cover_tail_include_prompts_and_overlap(self):
        import re
        class Processor:
            config=SimpleNamespace(max_len=100)
            def words_splitter(self,text):
                return [(m.group(),m.start(),m.end()) for m in re.finditer(r'\S+',text)]
            def prepare_inputs(self,texts,labels): return [['prompt']*10+texts[0]], [10]
            def transformer_tokenizer(self,inputs,**kwargs): return {'input_ids':[list(range(len(inputs[0])*2))]}
        words=[f'word{i}' for i in range(130)]
        windows=list(ner_windows(' '.join(words),Processor(),limit=64,overlap=5))
        self.assertTrue(all((len(w.split())+10)*2 <= 64 for w in windows))
        self.assertEqual(set(' '.join(windows).split()),set(words))
        self.assertTrue(all(set(a.split()) & set(b.split()) for a,b in zip(windows,windows[1:])))
        self.assertIn('word129',windows[-1])


if __name__=='__main__': unittest.main()
