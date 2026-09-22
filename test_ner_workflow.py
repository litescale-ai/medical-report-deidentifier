"""Use real documents to prove the shared NER export and review boundary."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pymupdf

from utils.clinical_packet import Section, extract_case, prepare_packet, packet_findings, PacketReviewRequired
from utils.document_formats import extract_document
from utils.ner_workflow import prepare_outputs, approve_outputs
from utils.reidentification import restore_batch
from test_clinical_packet import Detector, person


class NerWorkflowTest(unittest.IsolatedAsyncioTestCase):
    async def test_wrapped_pdf_identifier_with_touching_font_boxes_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory)
            source = home / '2025_Grade 3 Term 4_Report.pdf'
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((40, 60), 'Clinical skills', fontsize=9)
                page.insert_text((40, 72.2), 'Alex', fontsize=9)
                page.insert_text((40, 84.4), 'Example', fontsize=9)
                pdf.save(source)
            result = await prepare_outputs([source], output=home / 'case', secure=home / 'secure',
                                           detector=Detector([person('Alex\nExample')]), formats=['pdf'])
            self.assertFalse(result['failed'])
            text = '\n'.join(value for _, value in extract_document(result['pdfs'][0]['path']))
            self.assertNotIn('Alex', text)
            self.assertNotIn('Example', text)
            self.assertIn('Clinical skills', text)

    async def test_each_format_has_its_own_destination_and_manifest_entry(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            root = home / 'Reports'
            root.mkdir()
            source = root / 'Report.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((40, 60), 'Alex Example. Dose 5 mg.')
                pdf.save(source)
            folders = {'markdown': str(home / 'Reviewed text'), 'pdf': str(home / 'Reviewed PDFs')}
            result = await prepare_outputs([source], source_root=root, output=home / 'case', secure=home / 'secure',
                detector=Detector([person('Alex Example')]), formats=['markdown', 'pdf'], export_folders=folders)
            exports = approve_outputs(result, result['markdown_documents'])
            self.assertEqual(exports, [str(home / 'Reviewed text/document-001.md'),
                                       str(home / 'Reviewed PDFs/document-001.pdf')])
            self.assertEqual(result['reviewed_folders'], folders)
            private = json.loads((home / 'case/PRIVATE.json').read_text())
            self.assertEqual(private['export_folders'], folders)
            self.assertEqual(private['export_history'][-1]['folders'], folders)
            for kind, folder in folders.items():
                paths = [path for path in exports if Path(path).parent == Path(folder)]
                restored = restore_batch(paths, root=folder, output=home / ('restored-' + kind), secure=home / 'secure')
                self.assertFalse(restored['failed'])
                self.assertIn('Alex Example', '\n'.join(text for _, text in extract_document(restored['completed'][0])))
            # A failure at the second destination must not leave a partial new export.
            blocked = home / 'blocked'
            blocked.write_text('Existing user file')
            result['export_folders']['pdf'] = str(blocked / 'pdf')
            with self.assertRaises(OSError):
                approve_outputs(result, result['markdown_documents'])
            self.assertFalse(list((home / 'Reviewed text-2').glob('*.md')))
            self.assertTrue(all(Path(path).exists() for path in exports))
            self.assertEqual(blocked.read_text(), 'Existing user file')
            self.assertEqual(json.loads((home / 'case/PRIVATE.json').read_text())['reviewed_outputs'], exports)
            for invalid in [{'markdown': folders['markdown']}, {'markdown': '', 'pdf': folders['pdf']},
                            {'markdown': str(root), 'pdf': folders['pdf']},
                            {'markdown': folders['pdf'], 'pdf': folders['pdf']}]:
                with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                    await prepare_outputs([source], source_root=root, output=home / 'invalid', secure=home / 'secure',
                        detector=Detector([]), formats=['markdown', 'pdf'], export_folders=invalid)

    async def test_flagged_text_is_specific_and_override_is_explicit_and_audited(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            source = home / 'source.txt'
            source.write_text('Alex Example. Dose 5 mg.')
            result = await prepare_outputs([source], output=home / 'case', secure=home / 'secure',
                                           export_folder=home / 'source-redacted',
                                           detector=Detector([person('Alex Example')]), formats=['markdown'])
            edited = result['markdown'] + '\n## Document 1, section 1\nALEX   EXAMPLE called 0215550123.\n'
            findings = packet_findings(home / 'case', edited)
            self.assertEqual({item['type'] for item in findings}, {'PERSON', 'PHONE'})
            person_finding = next(item for item in findings if item['type'] == 'PERSON')
            self.assertEqual(person_finding['text'], 'ALEX   EXAMPLE')
            self.assertEqual(person_finding['location'], 'Document 1, section 1')
            self.assertEqual(edited.splitlines()[person_finding['line'] - 1], 'ALEX   EXAMPLE called 0215550123.')
            self.assertEqual(person_finding['context'], 'ALEX EXAMPLE called 0215550123.')
            with self.assertRaises(PacketReviewRequired) as raised:
                approve_outputs(result, edited)
            self.assertEqual(raised.exception.findings, findings)
            self.assertFalse(list((home / 'case').glob('REVIEWED-*')))
            exports = approve_outputs(result, edited, override=True, review_note='Reviewed intentionally retained text.')
            self.assertIn('ALEX   EXAMPLE', Path(exports[0]).read_text())
            manifest = json.loads((home / 'case/PRIVATE.json').read_text())
            self.assertEqual(manifest['verification'], 'user_override')
            audit = manifest['reviews'][-1]
            self.assertEqual(audit['accepted_findings'], findings)
            self.assertEqual(audit['note'], 'Reviewed intentionally retained text.')
            self.assertEqual(audit['verification'], 'user_override')
            # A previous override never silently authorises another save.
            with self.assertRaises(PacketReviewRequired):
                approve_outputs(result, edited + ' New edit.')
            with self.assertRaisesRegex(ValueError, 'empty'):
                approve_outputs(result, '', override=True)

    async def test_pdf_name_is_removed_when_font_boxes_touch_adjacent_lines(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            source = home / 'touching.pdf'
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((40, 60), 'Alex Example', fontsize=9)
                page.insert_text((40, 72.2), 'Clinical text stays.', fontsize=9)
                pdf.save(source)
            result = await prepare_outputs([source], secure=home / "secure", output=home / 'case', detector=Detector([person('Alex Example')]), formats=['pdf'])
            self.assertFalse(result['failed'])
            text = '\n'.join(value for _, value in extract_document(result['pdfs'][0]['path']))
            self.assertNotIn('Alex Example', text)
            self.assertIn('Clinical text stays.', text)

    async def test_both_exports_preserve_pdf_content_and_keep_terms_with_private_approval(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            files = []
            for number in range(2):
                source = home / f'private-name-{number}.pdf'
                with pymupdf.open() as pdf:
                    for page_number in range(2):
                        page = pdf.new_page()
                        page.insert_text((40, 60), 'Alex Example met Jamie Review.')
                        page.insert_text((40, 90), 'Practice No. 1270753')
                        page.insert_text((40, 140), 'Beery VMI score 85. Dose 5 mg.')
                        page.insert_text((40, 180), 'VMI')
                    pdf.save(source)
                files.append(source)
            before = [source.read_bytes() for source in files]
            detector = Detector([person('Alex Example'), person('VMI')])
            progress = []
            result = await prepare_outputs(files, secure=home / "secure", output=home / 'case', detector=detector,
                export_folder=home / 'reports-redacted',
                formats=['markdown', 'pdf'], keep_terms=['Beery VMI'], remove_terms=['Jamie Review'], progress=progress.append)
            self.assertFalse(result['failed'])
            self.assertEqual(len(result['pdfs']), 2)
            self.assertEqual(result['stats']['pages'], 4)
            self.assertTrue(any(value.startswith('Identifying section 4 / 4: private-name-1.pdf') for value in progress))
            for pdf in result['pdfs']:
                with pymupdf.open(pdf['path']) as actual, pymupdf.open(files[pdf['document'] - 1]) as original:
                    for page, old in zip(actual, original):
                        text = page.get_text()
                        self.assertNotIn('Alex Example', text)
                        self.assertNotIn('Jamie Review', text)
                        self.assertNotIn('1270753', text)
                        self.assertIn('Beery VMI score 85. Dose 5 mg.', text)
                        # Clinical line pixels are preserved, not reconstructed from OCR.
                        clip = pymupdf.Rect(35, 125, 350, 148)
                        self.assertEqual(page.get_pixmap(clip=clip).samples, old.get_pixmap(clip=clip).samples)
            self.assertIn('Beery VMI', result['markdown'])
            exports = approve_outputs(result, result['markdown_documents'])
            self.assertEqual(len(exports), 4)
            self.assertIn('**USER REVIEWED.**', Path(exports[0]).read_text())
            manifest = json.loads((home / 'case/PRIVATE.json').read_text())
            self.assertEqual(manifest['verification'], 'user_approved')
            self.assertEqual(manifest['remove_terms'], ['Jamie Review'])
            self.assertEqual([source.read_bytes() for source in files], before)
            for path in exports:
                self.assertEqual(Path(path).stat().st_mode & 0o777, 0o600)
            self.assertEqual([Path(path).name for path in exports],
                             ['document-001.md', 'document-002.md', 'document-001.pdf', 'document-002.pdf'])
            restored = restore_batch(exports, root=home / 'reports-redacted', output=home / 'restored', secure=home / 'secure')
            self.assertFalse(restored['failed'])
            self.assertEqual(len(restored['completed']), 4)
            for path in restored['completed']:
                text = '\n'.join(value for _, value in extract_document(path))
                self.assertIn('Alex Example', text)
                self.assertNotIn('Jamie Review', text)
                self.assertNotIn('1270753', text)
            with self.assertRaisesRegex(ValueError, 'Known identifying'):
                approve_outputs(result, {**result['markdown_documents'], 'document-001.md': result['markdown_documents']['document-001.md'] + '\nAlex Example'})
            Path(result['pdfs'][0]['path']).write_bytes(b'edited')
            with self.assertRaisesRegex(ValueError, 'changed outside'):
                approve_outputs(result, result['markdown_documents'])

    async def test_sequential_exports_map_original_subfolders_and_preserve_previous_saves(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            root = home / 'Medical Reports'
            files = []
            for number, folder in enumerate(['Consultations', 'Assessments'], 1):
                source = root / folder / 'Report.pdf'
                source.parent.mkdir(parents=True)
                with pymupdf.open() as pdf:
                    pdf.new_page().insert_text((40, 60), f'Alex Example. Dose {number * 5} mg.')
                    pdf.save(source)
                files.append(source)
            originals = [source.read_bytes() for source in files]
            def extract_working_copies(paths):
                self.assertEqual([path.name for path in paths], ['document-001.pdf', 'document-002.pdf'])
                self.assertEqual([path.read_bytes() for path in paths], originals)
                self.assertTrue(all(path.stat().st_mode & 0o777 == 0o600 for path in paths))
                return extract_case(paths)

            with patch('utils.ner_workflow.extract_case', side_effect=extract_working_copies):
                result = await prepare_outputs(files, source_root=root, output=home / 'case', secure=home / 'secure',
                                               detector=Detector([person('Alex Example')]), formats=['markdown', 'pdf'])
            expected_folder = home / 'Medical Reports-redacted'
            self.assertFalse(expected_folder.exists())
            exports = approve_outputs(result, result['markdown_documents'])
            self.assertEqual([str(Path(path).relative_to(expected_folder)) for path in exports],
                             ['document-001.md', 'document-002.md', 'document-001.pdf', 'document-002.pdf'])
            self.assertFalse((expected_folder / 'PRIVATE.json').exists())
            for number, path in enumerate(exports[:2], 1):
                text = Path(path).read_text()
                self.assertIn(f'## document-{number:03d}, page 1', text)
                self.assertIn(f'Dose {number * 5} mg.', text)
                self.assertNotIn(f'Dose {(3 - number) * 5} mg.', text)
                self.assertNotIn('Report.pdf', text)
            private = json.loads((home / 'case/PRIVATE.json').read_text())
            self.assertEqual([entry['original_name'] for entry in private['filename_mapping']],
                             ['Consultations/Report.pdf', 'Assessments/Report.pdf'])
            for number, entry in enumerate(private['filename_mapping'], 1):
                self.assertEqual(entry['new_name'], f'document-{number:03d}.pdf')
                self.assertEqual(entry['reviewed_outputs'], {
                    'markdown': str(expected_folder / f'document-{number:03d}.md'),
                    'pdf': str(expected_folder / f'document-{number:03d}.pdf')})
                self.assertTrue((home / 'case' / f'REVIEW_REQUIRED-document-{number:03d}.md').exists())
            before = [Path(path).read_bytes() for path in exports]
            # Even a user-edited previous export is retained when saving again.
            Path(exports[0]).write_text('User changes')
            second = approve_outputs(result, result['markdown_documents'])
            self.assertTrue(all(Path(path).is_relative_to(home / 'Medical Reports-redacted-2') for path in second))
            self.assertEqual([Path(path).read_bytes() for path in second], before)
            self.assertEqual(Path(exports[0]).read_text(), 'User changes')
            self.assertEqual([source.read_bytes() for source in files], originals)
            manifest = json.loads((home / 'case/PRIVATE.json').read_text())
            self.assertEqual(len(manifest['export_history']), 2)
            self.assertEqual(manifest['reviewed_outputs'], second)

    async def test_each_markdown_is_reviewed_and_edits_stay_with_its_source(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory)
            files = [home / 'private-one.txt', home / 'private-two.txt']
            for number, source in enumerate(files, 1):
                source.write_text(f'Alex Example. Dose {number * 5} mg.')
            result = await prepare_outputs(files, output=home / 'case', secure=home / 'secure',
                export_folder=home / 'exports', detector=Detector([person('Alex Example')]), formats=['markdown'])
            edited = dict(result['markdown_documents'])
            with self.assertRaisesRegex(ValueError, 'one Markdown file per source'):
                approve_outputs(result, {'document-001.md': edited['document-001.md']})
            edited['document-002.md'] += '\nAlex Example\n'
            with self.assertRaises(PacketReviewRequired):
                approve_outputs(result, edited)
            self.assertFalse((home / 'exports').exists())
            edited['document-002.md'] = result['markdown_documents']['document-002.md'] + '\nReviewed second document.\n'
            exports = approve_outputs(result, edited)
            self.assertNotIn('Reviewed second document.', Path(exports[0]).read_text())
            self.assertIn('Reviewed second document.', Path(exports[1]).read_text())
            edited['document-001.md'] += '\nAlex Example\n'
            approve_outputs(result, edited, override=True)
            manifest = json.loads((home / 'case/PRIVATE.json').read_text())
            self.assertEqual(manifest['verification'], 'user_override')
            self.assertEqual(len(manifest['reviewed_markdown_sha256']), 2)

    async def test_pdf_export_is_withheld_when_an_identifier_survives(self):
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            source = home / 'source.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((40, 60), 'Alex Example. Dose 5 mg.')
                pdf.save(source)
            def faulty_export(source, target, *args, **kwargs):
                Path(target).write_bytes(Path(source).read_bytes())
            with patch('utils.ner_workflow.deidentify_pdf', faulty_export):
                result = await prepare_outputs([source], secure=home / "secure", output=home / 'case', detector=Detector([person('Alex Example')]),
                                                formats=['pdf'])
            self.assertFalse(result['pdfs'])
            self.assertIn('Known identifiers remain', result['failed']['source.pdf'])
            self.assertIn('Alex Example', result['failed']['source.pdf'])
            self.assertEqual(result['failure_details']['source.pdf'][0]['location'], 'source.pdf, page 1')
            self.assertEqual(result['failure_details']['source.pdf'][0]['type'], 'PERSON')
            self.assertFalse(list((home / 'case').glob('*.pdf')))

    async def test_manual_terms_cannot_be_silently_exempted_or_hallucinated(self):
        for value in ['VMI', 'Unknown']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                await prepare_packet([Section('page 1', 'Beery VMI')], Detector([]),
                                     keep_terms=['Beery VMI'], remove_terms=[value])
        with self.assertRaisesRegex(ValueError, 'required removal'):
            await prepare_packet([Section('page 1', 'Practice No. 1270753')], Detector([]),
                                 keep_terms=['Practice No. 1270753'])

    async def test_markdown_all_formats_and_pdf_input_validation(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            home = Path(directory).resolve()
            files = []
            text = 'Alex Example. Dose 5 mg.'
            for suffix in ['txt', 'md', 'html']:
                file = home / f'source.{suffix}'
                file.write_text(text)
                files.append(file)
            doc = Document()
            doc.add_paragraph(text)
            doc.save(home / 'source.docx')
            book = Workbook()
            book.active['A1'] = text
            book.save(home / 'source.xlsx')
            slides = Presentation()
            slides.slides.add_slide(slides.slide_layouts[1]).shapes.title.text = text
            slides.save(home / 'source.pptx')
            files += [home / ('source.' + suffix) for suffix in ('docx', 'xlsx', 'pptx')]
            result = await prepare_outputs(files, secure=home / "secure", output=home / 'case', detector=Detector([person('Alex Example')]), formats=['markdown'])
            self.assertEqual(result['stats']['documents'], 6)
            self.assertNotIn('Alex Example', result['markdown'])
            self.assertEqual(result['markdown'].count('Dose 5 mg.'), 6)
            self.assertFalse(result['pdfs'])
            with self.assertRaisesRegex(ValueError, 'PDF inputs'):
                await prepare_outputs(files, secure=home / "secure", output=home / 'invalid', detector=Detector([]), formats=['pdf'])


if __name__ == '__main__':
    unittest.main()
