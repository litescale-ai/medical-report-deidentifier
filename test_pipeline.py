"""Focused local processing regressions; no model or server is required."""

import asyncio
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import httpx

from agents.cataloguer import catalogue_transcripts
from agents.deidentifier import discover_pii_entities
from agents.transcriber import transcribe_media, ExtractedTranscript


class LocalPipelineTest(unittest.IsolatedAsyncioTestCase):
    async def test_discovery_uses_source_spans_for_case_and_spacing_variants(self):
        from unittest.mock import AsyncMock
        from utils.batch import discover_names
        from utils.identifier_rules import replace_data
        response = {'entities': [{'name': 'South Africa', 'kind': 'LOCATION',
                                  'aliases': ['country of operation for address changes']}]}
        text = 'Jurisdiction: SOUTH  AFRICA. Office: South\nAfrica. Visit south africa.'
        with patch('utils.batch.generate_structured', AsyncMock(return_value=response)):
            entities = await discover_names(text, model='test', base_url='http://localhost')
        entity = entities[0]
        spans = [entity['canonical_name'], *entity['variations']]
        self.assertEqual(set(spans), {'SOUTH  AFRICA', 'South\nAfrica', 'south africa'})
        self.assertTrue(all(span in text for span in spans))
        self.assertEqual(replace_data(text, dict.fromkeys(spans, '[LOCATION]')),
                         'Jurisdiction: [LOCATION]. Office: [LOCATION]. Visit [LOCATION].')

    async def test_discovery_still_rejects_invented_and_partial_identifiers(self):
        from unittest.mock import AsyncMock
        from utils.batch import discover_names
        for name, text in (('South Africa', 'Jurisdiction: France.'), ('Ann', 'Annual review.'),
                           ('123', 'Reference: 1234'), ('', 'No names.')):
            response = {'entities': [{'name': name, 'kind': 'IDENTIFIER', 'aliases': []}]}
            with self.subTest(name=name), patch('utils.batch.generate_structured', AsyncMock(return_value=response)):
                with self.assertRaisesRegex(ValueError, 'absent|empty'):
                    await discover_names(text, model='test', base_url='http://localhost')

    async def test_timed_out_discovery_splits_once_and_preserves_all_sections(self):
        from unittest.mock import AsyncMock
        from utils.batch import chunks, discover_names
        text = ('Clinical review. ' * 220) + ' South Africa.'
        response = {'entities': [{'name': 'South Africa', 'kind': 'LOCATION', 'aliases': []}]}
        parts = list(chunks(text, limit=3000))
        replies = [TimeoutError('deadline'), {'entities': []}, response]
        metrics = lambda measured: None
        with patch('utils.batch.generate_structured', AsyncMock(side_effect=replies)) as generate:
            entities = await discover_names(text, model='test', base_url='http://localhost', metrics_callback=metrics)
        self.assertEqual([call.args[0] for call in generate.await_args_list], [text, *parts])
        self.assertTrue(all(call.kwargs['metrics_callback'] is metrics for call in generate.await_args_list))
        self.assertEqual(entities[0]['canonical_name'], 'South Africa')
        with patch('utils.batch.generate_structured', AsyncMock(side_effect=TimeoutError('deadline'))) as generate:
            with self.assertRaises(TimeoutError):
                await discover_names(text, model='test', base_url='http://localhost')
        self.assertEqual(generate.await_count, 2)  # Original and first smaller piece; no endless retries.

    async def test_default_model_and_explicit_overrides_reach_ollama(self):
        import os
        for environment, selected, expected in (
            ({}, None, 'qwen3.5:2b'),
            ({'OLLAMA_MODEL': 'gemma4:e2b'}, None, 'gemma4:e2b'),
            ({'OLLAMA_MODEL': 'gemma4:e2b'}, 'gemma4:e4b', 'gemma4:e4b'),
        ):
            with self.subTest(expected=expected):
                def respond(request):
                    self.assertEqual(json.loads(request.content)['model'], expected)
                    return httpx.Response(200, json={'done': True, 'done_reason': 'stop',
                        'message': {'content': '{"entities": []}'}})
                client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
                with patch.dict(os.environ, environment, clear=True), patch(
                    'utils.agent_config.httpx.AsyncClient', return_value=client
                ):
                    await discover_pii_entities({}, backend='ollama', ollama_model=selected)

    async def test_native_ollama_metrics_use_nanoseconds_and_missing_is_unknown(self):
        from utils.batch import discover_names
        for telemetry, expected in [({'prompt_eval_count': 64, 'eval_count': 48, 'eval_duration': 2_000_000_000}, 24), ({}, None)]:
            def respond(request):
                return httpx.Response(200, json={'done': True, 'done_reason': 'stop',
                    'message': {'content': '{"entities": []}'}, **telemetry})
            client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
            metrics = []
            with patch('utils.agent_config.httpx.AsyncClient', return_value=client):
                await discover_names('No names.', model='test', base_url='http://localhost:11434', metrics_callback=metrics.append)
            self.assertEqual(len(metrics), 1)
            self.assertEqual(metrics[0]['tokens_per_second'], expected)
            self.assertGreater(metrics[0]['request_seconds'], 0)

    async def test_total_deadline_cancels_a_stalled_request(self):
        cancelled = asyncio.Event()
        async def respond(request):
            try:
                await asyncio.Event().wait()
            finally:
                cancelled.set()
        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with patch("utils.agent_config.httpx.AsyncClient", return_value=client), patch.dict(
            "os.environ", {"MODEL_TIMEOUT_SECONDS": "0.02"}
        ):
            with self.assertRaisesRegex(TimeoutError, "deadline"):
                await discover_pii_entities({}, backend="ollama")
        self.assertTrue(cancelled.is_set())

    async def test_docx_extraction_preserves_body_and_tables(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "record.docx"
            doc = Document()
            doc.add_paragraph("Patient: Alex Example")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Doctor: Robin Test"
            doc.save(path)
            with patch("agents.transcriber.Agent", side_effect=AssertionError("Model used to copy DOCX")):
                result = await transcribe_media(str(path), backend="ollama")
            self.assertEqual(result["items"][0]["content"], "Patient: Alex Example\nDoctor: Robin Test")

    async def test_unavailable_server_does_not_fall_back_to_cloud(self):
        def respond(request):
            raise httpx.ConnectError("offline", request=request)
        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with patch("utils.agent_config.httpx.AsyncClient", return_value=client), patch(
            "utils.agent_config.Agent", side_effect=AssertionError("Cloud fallback")
        ):
            with self.assertRaisesRegex(RuntimeError, "Cannot reach Ollama"):
                await discover_pii_entities({}, backend="ollama")

    async def test_text_extraction_preserves_content_without_model(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "record.txt"
            text = "Patient: Alex Example.\n\nClinician: Dr. Robin Test.\n"
            path.write_text(text)
            with patch("agents.transcriber.Agent", side_effect=AssertionError("Model used to copy text")):
                result = await transcribe_media(str(path), backend="ollama")
            ExtractedTranscript.model_validate(result)
            self.assertEqual(result["items"][0]["content"], text)
            self.assertEqual(result["filename"], "record.txt")

    async def test_local_stages_use_native_schema_and_selected_endpoint(self):
        chronology = {
            "patient_summary": "Alex Example attended.", "categories_found": ["Intake"],
            "chronology": [{"timestamp": "line 1", "category": "Intake",
                            "source_file": "record.txt", "speaker": "Document",
                            "event_details": "Patient: Alex Example."}],
        }
        entities = {"entities": [{"canonical_name": "Alex Example", "entity_type": "PATIENT",
                                  "relationship_context": "Subject", "variations": ["Alex Example"]}]}
        requests = []

        def respond(request):
            requests.append(request)
            payload = json.loads(request.content)
            schema = payload["format"]
            data = chronology if schema["title"] == "UnifiedChronology" else entities
            return httpx.Response(200, json={"done": True, "done_reason": "stop",
                                           "message": {"content": json.dumps(data)}})

        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with patch("utils.agent_config.httpx.AsyncClient", return_value=client):
            result = await catalogue_transcripts([], backend="ollama", ollama_base_url="http://local.test:1234/v1")
        self.assertEqual(result, chronology)
        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with patch("utils.agent_config.httpx.AsyncClient", return_value=client):
            result = await discover_pii_entities(chronology, backend="ollama", ollama_base_url="http://local.test:1234/v1")
        self.assertEqual(result, entities["entities"])
        self.assertEqual(len(requests), 2)
        for request in requests:
            self.assertEqual(str(request.url), "http://local.test:1234/api/chat")
            payload = json.loads(request.content)
            self.assertFalse(payload["think"])
            self.assertNotIn("tools", payload)
            self.assertIn("$defs", payload["format"])

    async def test_bad_outputs_fail_once_without_silent_success(self):
        for response in (
            {"done": True, "done_reason": "stop", "message": {"content": "not json"}},
            {"done": True, "done_reason": "stop", "message": {"content": '{"entities":[{}]}'}},
            {"done": True, "done_reason": "length", "message": {"content": '{"entities":[]}'}},
            {"done": False, "message": {"content": '{"entities":[]}'}},
        ):
            with self.subTest(response=response):
                calls = []
                def respond(request):
                    calls.append(request)
                    return httpx.Response(200, json=response)
                client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
                with patch("utils.agent_config.httpx.AsyncClient", return_value=client):
                    with self.assertRaisesRegex(ValueError, "Ollama"):
                        await discover_pii_entities({}, backend="ollama")
                self.assertEqual(len(calls), 1)

    async def test_network_timeout_is_actionable_and_not_retried(self):
        async def respond(request):
            raise httpx.ReadTimeout("test", request=request)
        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with patch("utils.agent_config.httpx.AsyncClient", return_value=client):
            with self.assertRaisesRegex(TimeoutError, "Ollama"):
                await discover_pii_entities({}, backend="ollama")


class PdfRedactionTest(unittest.TestCase):
    def test_pdf_overlapping_aliases_and_no_match_do_not_trigger_ocr(self):
        import pymupdf
        from utils.document_editor import deidentify_pdf
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.pdf", Path(directory) / "output.pdf"
            with pymupdf.open() as doc:
                doc.new_page().insert_text((40, 60), "Alex Example attended.")
                doc.save(source)
            with patch("utils.document_editor._ocr_pdf", side_effect=AssertionError("Unnecessary OCR")):
                self.assertTrue(deidentify_pdf(str(source), str(output), {"Alex Example": "PATIENT_TEST", "Alex": "PATIENT_TEST"}))
                with pymupdf.open(output) as doc:
                    self.assertEqual(doc[0].get_text().count("PATIENT_TEST"), 1)
                output.unlink()
                self.assertTrue(deidentify_pdf(str(source), str(output), {}))

    def test_scanned_pdf_ocr_failure_does_not_export_original(self):
        import pymupdf
        from utils.document_editor import deidentify_pdf
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "scan.pdf", Path(directory) / "output.pdf"
            with pymupdf.open() as original:
                page = original.new_page()
                page.insert_text((40, 60), "Alex Example")
                pixmap = page.get_pixmap()
                with pymupdf.open() as scan:
                    scan.new_page().insert_image(page.rect, pixmap=pixmap)
                    scan.save(source)
            with patch("utils.document_editor._ocr_pdf", side_effect=RuntimeError("OCR unavailable")):
                with self.assertRaisesRegex(RuntimeError, "OCR unavailable"):
                    deidentify_pdf(str(source), str(output), {"Alex Example": "PATIENT_TEST"})
            self.assertFalse(output.exists())

    def test_styles_are_extracted_once_per_page_without_changing_redaction(self):
        import pymupdf
        from utils.document_editor import _redact_pdf
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.pdf"
            output = Path(directory) / "output.pdf"
            with pymupdf.open() as doc:
                for _ in range(2):
                    page = doc.new_page()
                    for line in range(3):
                        page.insert_text((40, 40 + line * 30), "Alex Example", fontsize=11)
                doc.save(source)
            extracted = []
            original = pymupdf.Page.get_text
            def get_text(page, *args, **kwargs):
                if args and args[0] == "dict":
                    extracted.append(page.number)
                return original(page, *args, **kwargs)
            with patch.object(pymupdf.Page, "get_text", get_text):
                count = _redact_pdf(str(source), str(output), {"Alex Example": "PATIENT_ABC"})
            self.assertEqual(count, 6)
            self.assertEqual(extracted, [0, 1])
            with pymupdf.open(output) as doc:
                for xref in range(1, doc.xref_length()):
                    if doc.xref_is_stream(xref):
                        stream = doc.xref_stream(xref)
                        self.assertNotIn(b"Alex Example", stream)
                        self.assertNotIn(b"416c6578204578616d706c65", stream.lower())
                for page in doc:
                    self.assertNotIn("Alex Example", page.get_text())
                    self.assertEqual(page.get_text().count("PATIENT_ABC"), 3)


if __name__ == "__main__":
    unittest.main()
