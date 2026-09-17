"""Local text extraction shared by transcription and document editing."""

import re
from html import escape, unescape
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".log", ".xml", ".yaml", ".yml", ".html", ".htm"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {".xlsx", ".docx", ".pdf", ".pptx"}


def read_text(path):
    """Accept UTF-8 (including BOM) and BOM-marked UTF-16; reject binary input."""
    data = Path(path).read_bytes()
    encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    text = data.decode(encoding)
    if "\x00" in text:
        raise ValueError("The file contains binary data, not supported plain text.")
    return text


def replace_strings(strings, replacements):
    """Replace longest matches once across formatted runs, keeping other runs intact."""
    keys = sorted((key for key in replacements if key), key=len, reverse=True)
    result = list(strings)
    if not keys:
        return result
    text = "".join(strings)
    positions = [(ri, ci) for ri, value in enumerate(strings) for ci in range(len(value))]
    # Right to left keeps original offsets valid, even with adjacent matches.
    for match in reversed(list(re.finditer("|".join(map(re.escape, keys)), text))):
        first, start = positions[match.start()]
        last, end = positions[match.end() - 1]
        suffix = result[last][end + 1:]
        if first == last:
            result[first] = result[first][:start] + replacements[match.group()] + suffix
        else:
            result[first] = result[first][:start] + replacements[match.group()]
            for index in range(first + 1, last):
                result[index] = ""
            result[last] = suffix
    return result


def docx_paragraphs(container):
    """Walk body/header/footer paragraphs and nested tables without duplicate cells."""
    yield from container.paragraphs
    seen = set()
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc)
                    yield from docx_paragraphs(cell)


def all_docx_paragraphs(document):
    yield from docx_paragraphs(document)
    for section in document.sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            if not part.is_linked_to_previous:
                yield from docx_paragraphs(part)


def slide_paragraphs(shapes):
    """Include grouped shapes and table cells as well as ordinary slide text."""
    for shape in shapes:
        if shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs
        if hasattr(shape, "shapes"):
            yield from slide_paragraphs(shape.shapes)


def presentation_paragraphs(presentation):
    for number, slide in enumerate(presentation.slides, 1):
        yield f"slide {number}", list(slide_paragraphs(slide.shapes))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                yield f"slide {number} notes", list(frame.paragraphs)


class HtmlText(HTMLParser):
    """Keep markup while replacing decoded text, including names split by inline tags."""
    def __init__(self, source):
        super().__init__(convert_charrefs=True)
        self.raw_tag = None
        self.parts = []  # (kind, value); boundary text is for matching only
        self.feed(source)
        self.close()

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.raw_tag = tag
        if tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "section"}:
            self.parts.append(("boundary", "\n"))
        self.parts.append(("markup", self.get_starttag_text()))

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)

    def handle_endtag(self, tag):
        if tag == self.raw_tag:
            self.raw_tag = None
        self.parts.append(("markup", f"</{tag}>"))
        if tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "section"}:
            self.parts.append(("boundary", "\n"))

    def handle_data(self, data):
        self.parts.append(("rawtext" if self.raw_tag else "text", data))

    def handle_decl(self, decl):
        self.parts.append(("markup", f"<!{decl}>"))

    def handle_comment(self, data):
        self.parts.append(("markup", f"<!--{data}-->"))

    def text(self):
        # Attributes/comments can contain PII too; include them in discovery.
        return "".join(value for kind, value in self.parts if kind != "markup") + "\n" + "\n".join(
            unescape(value) for kind, value in self.parts if kind == "markup")

    def replace(self, replacements):
        strings = [value if kind != "markup" else "" for kind, value in self.parts]
        changed = replace_strings(strings, replacements)
        def attribute(match):
            value = match.group(2)
            quote = value[0] if value[0] in {"\"", "'"} else ""
            decoded = unescape(value[1:-1] if quote else value)
            updated = replace_strings([decoded], replacements)[0]
            if updated == decoded:
                return match.group()
            return match.group(1) + '"' + escape(updated, quote=True) + '"'

        result = []
        for index, (kind, value) in enumerate(self.parts):
            if kind == "text":
                result.append(escape(changed[index], quote=False))
            elif kind == "rawtext":
                result.append(changed[index])
            elif kind == "markup":
                value = re.sub(r"([\w:-]+\s*=\s*)(\"[^\"]*\"|'[^']*'|[^\s>]+)", attribute, value)
                result.append(replace_strings([value], replacements)[0])
        return "".join(result)


def extract_document(path):
    """Return (location, text) sections, or None for media needing model transcription."""
    ext = Path(path).suffix.lower()
    if ext in {".doc", ".xls", ".ppt"}:
        raise ValueError(f"Convert {ext} to {ext}x before processing.")
    if ext in TEXT_EXTENSIONS:
        text = read_text(path)
        if ext in {".html", ".htm"}:
            text = HtmlText(text).text()
        return [("line 1", text)]
    if ext == ".docx":
        from docx import Document
        return [("document", "\n".join(p.text for p in all_docx_paragraphs(Document(path)) if p.text))]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        try:
            sections = []
            for sheet in workbook:
                values = []
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            values.append(f"{cell.coordinate}: {cell.value}")
                        if cell.comment:
                            values.append(f"{cell.coordinate} comment ({cell.comment.author}): {cell.comment.text}")
                        if cell.hyperlink and cell.hyperlink.target:
                            values.append(f"{cell.coordinate} link: {cell.hyperlink.target}")
                sections.append((f"sheet {sheet.title}", "\n".join(values)))
            return sections
        finally:
            workbook.close()
    if ext == ".pptx":
        from pptx import Presentation
        return [(location, "\n".join(p.text for p in paragraphs if p.text))
                for location, paragraphs in presentation_paragraphs(Presentation(path))]
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as document:
            sections = []
            for page in document:
                text = page.get_text()
                # OCR each image-only page locally, including mixed scanned/text PDFs.
                if not text.strip() and page.get_images():
                    text = page.get_text(textpage=page.get_textpage_ocr(language="eng", full=True))
                sections.append((f"page {page.number + 1}", text))
            return sections
    return None
