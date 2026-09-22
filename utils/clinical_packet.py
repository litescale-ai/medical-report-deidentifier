"""Experimental, local-only detection and Markdown preparation for one patient.

Always produces a review draft, never a claim of anonymisation. Identity maps
are case-local; no aliases or clinical relationships are inferred by NER.
"""
from collections import Counter
from dataclasses import dataclass
import json
import math
import os
from pathlib import Path
import re
from time import perf_counter

from utils.document_formats import extract_document
from utils.identifier_rules import identifier_replacements, replace_data, replacement_pattern

MODEL = 'knowledgator/gliner-pii-base-v1.0'
REVISION = '61726e0ad791dcab3e29339bbec3ad42ded65641'
# Match the existing deidentifier's exclusion of role words from aliases.
GENERIC_ROLES = {'patient', 'the patient', 'doctor', 'the doctor', 'dr', 'dr.', 'mr', 'mr.',
                 'mrs', 'mrs.', 'ms', 'ms.', 'clinician', 'relative', 'mother', 'father'}
LABELS = {
    'name': 'PERSON', 'organization': 'ORGANIZATION', 'school': 'ORGANIZATION',
    'organization medical facility': 'FACILITY', 'location address': 'LOCATION',
    'location city': 'LOCATION', 'location zip': 'LOCATION', 'dob': 'DOB',
    'healthcare number': 'IDENTIFIER', 'passport number': 'IDENTIFIER',
}
# Labelled IDs only: an unlabelled number could be a clinical measurement.
LABELLED_ID = re.compile(r'(?im)\b(?:SA\s+ID|ID\s*(?:number|no\.?)?|identity\s+number|'
                         r'patient\s+(?:number|no\.?)|file\s+(?:number|no\.?)|'
                         r'medical\s+record\s+(?:number|no\.?))[ \t]*[:#][ \t]*((?!\[)[^\s,;]+)')


@dataclass(frozen=True)
class Section:
    source: str
    text: str


def extract_case(files):
    """Use neutral provenance; keep DOCX table cells in rows instead of flattening them."""
    sections = []
    for number, filename in enumerate(files, 1):
        path = Path(filename)
        if path.suffix.lower() == '.docx':
            from docx import Document
            from docx.table import Table

            def blocks(container):
                for block in container.iter_inner_content():
                    if isinstance(block, Table):
                        for row in block.rows:
                            yield ' | '.join(' / '.join(blocks(cell)) for cell in row.cells)
                    else:
                        yield block.text
            doc = Document(path)
            parts = [('body', '\n'.join(blocks(doc)))]
            for index, section in enumerate(doc.sections, 1):
                for kind in ('header', 'first_page_header', 'even_page_header',
                             'footer', 'first_page_footer', 'even_page_footer'):
                    part = getattr(section, kind)
                    if not part.is_linked_to_previous:
                        parts.append((f'section {index} {kind}', '\n'.join(blocks(part))))
        else:
            parts = extract_document(path)
        if not parts or not any(text.strip() for _, text in parts):
            raise ValueError(f'Document {number} has no readable text; no packet written.')
        for index, (location, text) in enumerate(parts, 1):
            # Sheet titles and other source labels can contain identities.
            if path.suffix.lower() == '.pdf':
                source = f'Document {number}, page {index}'
            else:
                source = f'Document {number}, section {index}'
            if path.suffix.lower() == '.xlsx':
                text = location + '\n' + text
            sections.append(Section(source, text))
    if not sections:
        raise ValueError('Select at least one document for a single patient.')
    return sections


def rule_replacements(text):
    result = identifier_replacements(text)
    for match in LABELLED_ID.finditer(text):
        result[match[1]] = '[IDENTIFIER REMOVED]'
    return result


def ner_windows(text, processor, *, limit=384, overlap=24):
    """Count the actual tokenizer input, including label prompts; never truncate.

    Windows overlap in whole words to cover boundary-spanning names. The final
    span offsets are local to each window. One unusually long token fails closed.
    """
    words = list(processor.words_splitter(text))
    start = 0
    while start < len(words):
        end = min(start + limit, len(words))
        while True:
            values = [word for word, _, _ in words[start:end]]
            inputs, _ = processor.prepare_inputs([values], list(LABELS))
            ids = processor.transformer_tokenizer(inputs, is_split_into_words=True,
                                                   truncation=False, verbose=False)['input_ids'][0]
            if len(ids) <= limit and len(values) <= processor.config.max_len:
                break
            if end - start == 1:
                raise ValueError('One token exceeds the NER window; review the extraction.')
            end = start + max(1, (end - start) // 2)
        yield text[words[start][1]:words[end - 1][2]]
        if end == len(words):
            break
        start = max(start + 1, end - overlap)


class NerDetector:
    """Load the pinned local model once. Downloads contain model files, never reports."""
    def __init__(self, threshold=0.3):
        if not math.isfinite(threshold) or not 0 < threshold < 1:
            raise ValueError('NER threshold must be between zero and one.')
        os.environ.setdefault('HF_HUB_DISABLE_TELEMETRY', '1')
        from huggingface_hub import snapshot_download
        from gliner import GLiNER
        import torch
        # Fixed CPU execution makes this first comparison reproducible on this Mac.
        torch.set_num_threads(4)
        path = snapshot_download(MODEL, revision=REVISION,
                                 allow_patterns=['*.json', '*.model', 'pytorch_model.bin'])
        self.model = GLiNER.from_pretrained(path, local_files_only=True).eval()
        self.model.data_processor.transformer_tokenizer.model_max_length = 384
        self.threshold = threshold
        self.metadata = {'detector': MODEL, 'revision': REVISION, 'threshold': threshold,
                         'labels': list(LABELS), 'device': 'cpu', 'threads': 4}

    async def discover(self, text):
        entities = []
        for window in ner_windows(text, self.model.data_processor):
            for entity in self.model.predict_entities(window, list(LABELS), threshold=self.threshold):
                start, end = entity['start'], entity['end']
                if not 0 <= start < end <= len(window) or window[start:end] != entity['text']:
                    raise ValueError('NER returned an invalid span; no packet written.')
                entities.append({'canonical_name': entity['text'],
                                 'entity_type': LABELS[entity['label']], 'variations': []})
        return entities


class QwenDetector:
    """Use the existing detector for a like-for-like comparison; never cloud fallback."""
    def __init__(self):
        self.metadata = {'detector': 'qwen3.5:2b', 'base_url': 'http://127.0.0.1:11434'}

    async def discover(self, text):
        from utils.batch import chunks, discover_names
        entities = []
        for part in chunks(text):
            entities.extend(await discover_names(part, model=self.metadata['detector'],
                                                 base_url=self.metadata['base_url']))
        return entities


def reviewed_term_pattern(terms):
    """Match reviewed whole phrases with case/spacing tolerance, never substrings."""
    terms = sorted({' '.join(term.split()) for term in terms if term.strip()}, key=len, reverse=True)
    if not terms:
        return None
    return re.compile(replacement_pattern(terms).replace(r'\ ', r'\s+'), re.IGNORECASE)


def replace_with_review(text, replacements, keep_pattern):
    """Keep reviewed phrases intact; withhold conflicts that cross their boundaries.

    Protecting spans also prevents a short model alias from damaging a longer
    reviewed term. Return only values actually replaced for the private mapping.
    """
    kept = [match.span() for match in keep_pattern.finditer(text)] if keep_pattern else []
    used = set()
    if not replacements:
        return text, used
    pattern = re.compile(replacement_pattern(sorted(replacements, key=len, reverse=True)))

    def replace(match):
        if any(start <= match.start() and match.end() <= end for start, end in kept):
            return match.group()
        if any(start < match.end() and match.start() < end for start, end in kept):
            raise ValueError('An identifier crosses a reviewed keep-term boundary; review the conflicting detection.')
        used.add(match.group())
        return replacements[match.group()]

    return pattern.sub(replace, text), used


async def prepare_packet(sections, detector, *, keep_terms=(), remove_terms=(), on_progress=lambda completed, total: None):
    """Detect once per section, propagate exact variants within this case, then verify.

    Output verification checks known values only. It cannot detect a name missed
    everywhere by the detector. Independent review remains necessary.
    """
    started = perf_counter()
    keep_terms = list(dict.fromkeys(term.strip() for term in keep_terms if term.strip()))
    keep_pattern = reviewed_term_pattern(keep_terms)
    text = '\n\n'.join(section.text for section in sections)
    rules = rule_replacements(text)
    kept_spans = [match.span() for match in keep_pattern.finditer(text)] if keep_pattern else []
    if rules and any(start < match.end() and match.start() < end for start, end in kept_spans
                     for match in re.finditer(replacement_pattern(rules), text)):
        raise ValueError('A keep term contains a phone, registration, address or other required removal. Edit the keep list first.')
    for term in remove_terms:
        from utils.batch import source_variants
        variants = source_variants(text, term)
        if not variants:
            raise ValueError('Additional identifying text was not found in the source documents.')
        if any(start < match.end() and match.start() < end for start, end in kept_spans
               for match in re.finditer(replacement_pattern(variants), text)):
            raise ValueError('Additional identifying text conflicts with a keep term. Edit the keep list first.')
        rules.update({value: '[IDENTITY REMOVED]' for value in variants})
    entities = []
    on_progress(0, len(sections))
    for number, section in enumerate(sections, 1):
        entities.extend(await detector.discover(replace_data(section.text, rules)))
        on_progress(number, len(sections))
    replacements, mapping, counters, identities = {}, {}, Counter(), {}
    for entity in entities:
        name = entity['canonical_name'].strip()
        if not name:
            raise ValueError('Empty identifier; no packet written.')
        key = ' '.join(name.casefold().split())
        if key not in identities:
            kind = entity['entity_type']
            if kind not in {'PERSON', 'PATIENT', 'DOCTOR', 'RELATIVE', 'LOCATION', 'FACILITY',
                            'ORGANIZATION', 'DOB', 'IDENTIFIER', 'EMAIL'}:
                kind = 'IDENTIFIER'
            counters[kind] += 1
            identities[key] = f'[{kind}_{counters[kind]}]'
            mapping[identities[key]] = []
        token = identities[key]
        for value in [name, *entity.get('variations', [])]:
            if not value.strip() or value.strip().casefold() in GENERIC_ROLES:
                continue
            pattern = replacement_pattern([' '.join(value.split())]).replace(r'\ ', r'\s+')
            variants = dict.fromkeys(m.group() for m in re.finditer(pattern, text, re.IGNORECASE))
            if not variants:
                raise ValueError('Identifier absent from source; no packet written.')
            for variant in variants:
                if variant in replacements and replacements[variant] != token:
                    raise ValueError('Conflicting identity aliases; review before exporting.')
                replacements[variant] = token
                if variant not in mapping[token]:
                    mapping[token].append(variant)
    replacements.update(rules)
    cleaned, used = [], set()
    for section in sections:
        value, changed = replace_with_review(section.text, replacements, keep_pattern)
        cleaned.append(Section(section.source, value))
        used.update(changed)
    mapping = {token: [value for value in values if value in used and replacements.get(value) == token]
               for token, values in mapping.items() if any(value in used and replacements.get(value) == token for value in values)}
    replacements = {value: replacement for value, replacement in replacements.items() if value in used}
    cleaned_text = '\n\n'.join(s.text for s in cleaned)
    review_text = keep_pattern.sub(' ', cleaned_text) if keep_pattern else cleaned_text
    # Exclude generated markers when checking short names such as "Person".
    checked = re.sub(r'\[[A-Z]+(?:_\d+| REMOVED)\]', '', review_text)
    if rule_replacements(review_text) or (replacements and re.search(replacement_pattern(replacements), checked)):
        raise ValueError('Known identifiers remain; no packet written.')
    return {'markdown': render_packet(cleaned), 'mapping': mapping, 'keep_terms': keep_terms,
            'cleaned': [s.text for s in cleaned], 'replacements': replacements,
            'stats': {**detector.metadata, 'sections': len(sections), 'entities': len(mapping),
                      'rule_values': len(rules), 'kept_terms': len(keep_terms), 'seconds': round(perf_counter() - started, 3)}}


def render_packet(sections):
    """Render extracted sections with the same review notice and literal text boundaries."""
    parts = ['# Clinical evidence packet',
             '**REVIEW REQUIRED.** Automated pseudonymisation can miss identifiers or remove clinical text. '
             'Check against the originals before sharing. Keep placeholders unchanged when drafting. '
             'Different placeholders may refer to the same person; do not infer relationships.',
             'Text extraction may omit images, drawings, charts, annotations or table layout. '
             'Check these in the originals.']
    for section in sections:
        # A literal code block prevents source HTML/images/links from making network requests in preview.
        fence = '`' * max(3, 1 + max((len(m.group()) for m in re.finditer(r'`+', section.text)), default=0))
        parts.append(f'## {section.source}\n\n{fence}text\n{section.text}\n{fence}')
    return '\n\n'.join(parts) + '\n'


def write_packet(result, output, sources):
    """Create a new private case folder; never overwrite a previous packet or originals."""
    output = Path(output)
    output.mkdir(parents=True, mode=0o700, exist_ok=False)
    private = {'sources': [str(Path(p).resolve()) for p in sources],
               'identities': result['mapping'], 'replacements': result['replacements'],
               'keep_terms': result['keep_terms'], 'stats': result['stats']}
    if 'filename_mapping' in result:
        private['filename_mapping'] = result['filename_mapping']
    payloads = ({'REVIEW_REQUIRED-' + name: text for name, text in result['markdown_documents'].items()}
                if 'markdown_documents' in result else {'REVIEW_REQUIRED.md': result['markdown']})
    payloads['PRIVATE.json'] = json.dumps(private, indent=2)
    for name, content in payloads.items():
        fd = os.open(output / name, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, 'w') as stream:
            stream.write(content)


class PacketReviewRequired(ValueError):
    """Expose actionable findings without treating every detection as confirmed PII."""
    def __init__(self, findings):
        self.findings = findings
        values = list(dict.fromkeys(item['text'] for item in findings))
        super().__init__('Known identifying text remains: ' + ', '.join(repr(value) for value in values) +
                         '. Edit the flagged text or explicitly override after review.')


def packet_findings(output, text):
    """Report exact occurrences in the edited packet, with unchanged line offsets."""
    private = json.loads((Path(output) / 'PRIVATE.json').read_text())
    keep = reviewed_term_pattern(private['keep_terms'])
    def mask(match):
        return re.sub(r'[^\n]', ' ', match.group())
    checked = keep.sub(mask, text) if keep else text
    checked = re.sub(r'\[[A-Z]+(?:_\d+| REMOVED)\]', mask, checked)
    known = private.get('replacements', {})
    rules = rule_replacements(checked)
    values = {**known, **rules}
    pattern = reviewed_term_pattern(values)
    if pattern is None:
        return []
    normalize = lambda value: ' '.join(value.casefold().split())
    details = {normalize(value): (replacement, value in rules) for value, replacement in values.items()}
    headings = list(re.finditer(r'^## ((?:Document \d+|document-\d+)[^\n]*)$', text, re.MULTILINE))
    findings = []
    for match in pattern.finditer(checked):
        value = text[match.start():match.end()]
        if normalize(value) not in details:
            continue  # Masked keep terms must not create a match across removed text.
        replacement, from_rule = details[normalize(value)]
        kind = replacement.strip('[]').removesuffix(' REMOVED').rsplit('_', 1)[0]
        location = next((heading[1] for heading in reversed(headings) if heading.start() < match.start()), 'Packet text')
        line_start = text.rfind('\n', 0, match.start()) + 1
        line_end = text.find('\n', match.end())
        if line_end < 0:
            line_end = len(text)
        findings.append({'text': value, 'type': kind, 'location': location,
                         'line': text.count('\n', 0, match.start()) + 1,
                         'context': ' '.join(text[max(line_start, match.start() - 65):min(line_end, match.end() + 65)].split()),
                         'reason': 'Recognised by an identifier rule' if from_rule else 'Matches text identified earlier in this run'})
    return findings


def approve_packet(output, text, *, override=False, review_note=''):
    """Save explicitly reviewed text with an audit, without replacing earlier versions."""
    from hashlib import sha256
    from utils.batch import save_private
    from utils.batch_stats import utc_now
    output = Path(output)
    manifest = output / 'PRIVATE.json'
    private = json.loads(manifest.read_text())
    if not text.strip():
        raise ValueError('The reviewed packet cannot be empty.')
    findings = packet_findings(output, text)
    if findings and not override:
        raise PacketReviewRequired(findings)
    verification = 'user_override' if findings else 'user_approved'
    label = '**USER REVIEWED — FLAGGED TEXT RETAINED.**' if findings else '**USER REVIEWED.**'
    text = text.replace('**REVIEW REQUIRED.**', label, 1)
    fingerprint = sha256(text.encode('utf-8')).hexdigest()
    target = output / f'REVIEWED-{fingerprint[:12]}.md'
    if target.exists():
        if target.read_text(encoding='utf-8') != text:
            raise ValueError('An existing reviewed packet changed outside Guardian. Prepare a new packet.')
    else:
        fd = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            stream.write(text)
    private.setdefault('reviews', []).append({'file': target.name, 'sha256': fingerprint,
                                               'at': utc_now(), 'verification': verification,
                                               'accepted_findings': findings, 'note': review_note.strip()})
    save_private(manifest, private)
    return target
